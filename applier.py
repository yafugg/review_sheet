"""Apply approved corrections back to Excel or Word files."""

from pathlib import Path

from extractor import TextBlock
from detector import Correction


def apply_to_excel(path: Path, blocks: list[TextBlock], corrections: list[Correction]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)

    # Build a map from (sheet_title, coordinate) -> new text
    updates: dict[tuple[str, str], str] = {}
    for c in corrections:
        block = blocks[c.block_index]
        # location format: "シート「{title}」セル {coord}"
        try:
            parts = block.location.split("」セル ")
            sheet_title = parts[0].replace("シート「", "")
            coord = parts[1]
            updates[(sheet_title, coord)] = c.suggested
        except (IndexError, ValueError):
            print(f"[警告] セル位置を解析できませんでした: {block.location}")

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                key = (sheet.title, cell.coordinate)
                if key in updates:
                    cell.value = updates[key]

    wb.save(path)


def apply_to_word(path: Path, blocks: list[TextBlock], corrections: list[Correction]) -> None:
    from docx import Document

    doc = Document(path)

    # Build a map from block_index -> suggested text
    index_map: dict[int, str] = {c.block_index: c.suggested for c in corrections}

    para_idx = 0
    for para in doc.paragraphs:
        if para.text.strip():
            if para_idx in index_map:
                # Preserve runs structure by updating the first run and clearing others
                new_text = index_map[para_idx]
                if para.runs:
                    para.runs[0].text = new_text
                    for run in para.runs[1:]:
                        run.text = ""
                else:
                    para.text = new_text
            para_idx += 1

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    if para_idx in index_map:
                        # Replace text in first paragraph of cell
                        if cell.paragraphs:
                            p = cell.paragraphs[0]
                            new_text = index_map[para_idx]
                            if p.runs:
                                p.runs[0].text = new_text
                                for run in p.runs[1:]:
                                    run.text = ""
                            else:
                                p.text = new_text
                    para_idx += 1

    doc.save(path)


def apply_corrections(path: Path, blocks: list[TextBlock], corrections: list[Correction]) -> None:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xls", ".xlsm"):
        apply_to_excel(path, blocks, corrections)
    elif suffix in (".docx", ".doc"):
        apply_to_word(path, blocks, corrections)
    else:
        raise ValueError(f"未対応のファイル形式です: {suffix}")

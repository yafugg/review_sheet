"""Extract text with location info from Excel and Word files."""

from dataclasses import dataclass
from pathlib import Path
from typing import Iterator


@dataclass
class TextBlock:
    text: str
    location: str  # human-readable location description
    file_type: str  # "excel" or "word"


def extract_from_excel(path: Path) -> Iterator[TextBlock]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.strip():
                    yield TextBlock(
                        text=cell.value.strip(),
                        location=f"シート「{sheet.title}」セル {cell.coordinate}",
                        file_type="excel",
                    )


def extract_from_word(path: Path) -> Iterator[TextBlock]:
    from docx import Document

    doc = Document(path)

    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            yield TextBlock(
                text=para.text.strip(),
                location=f"段落 {i + 1}",
                file_type="word",
            )

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if cell.text.strip():
                    yield TextBlock(
                        text=cell.text.strip(),
                        location=f"表 {ti + 1} 行 {ri + 1} 列 {ci + 1}",
                        file_type="word",
                    )


def extract_text_blocks(path: Path) -> list[TextBlock]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xls", ".xlsm"):
        return list(extract_from_excel(path))
    elif suffix in (".docx", ".doc"):
        return list(extract_from_word(path))
    else:
        raise ValueError(f"未対応のファイル形式です: {suffix}")
